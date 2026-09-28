from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from statsmodels.stats.weightstats import CompareMeans, DescrStatsW

from fast_exp_analytics import export_ab_results_to_excel, run_ab_test, run_abc_test
from fast_exp_analytics.ab import (
    _prep_metric_pair_frames,
    calculate_base_exp_values,
    calculate_p_value,
)


def _metrics_config() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["Average", "average", "value", "value", "positive"],
            ["Ratio", "ratio", "numerator", "denominator", "positive"],
            ["Share", "share", "success", "eligible", "positive"],
            ["Median", "median", "median_value", "median_value", "positive"],
        ],
        index=["average", "ratio", "share", "median"],
        columns=["desc", "type", "num", "den", "direction"],
    )


def _experiment_data(groups: tuple[str, ...]) -> pd.DataFrame:
    rows = []
    for group_index, group in enumerate(groups):
        for value in range(1, 9):
            rows.append(
                {
                    "user_id": f"{group}-{value}",
                    "exp_group": group,
                    "value": value + group_index,
                    "numerator": value * (group_index + 1) + group_index**2,
                    "denominator": value + 2,
                    "success": int(value + group_index > 4),
                    "eligible": 1,
                    "median_value": value + group_index,
                }
            )
    return pd.DataFrame(rows)


def test_ab_returns_ci_for_every_metric_type():
    result = run_ab_test(
        _experiment_data(("A", "B")),
        _metrics_config(),
        "2026-01-01",
        "2026-01-07",
        ci_bootstrap_resamples=500,
        ci_random_state=1,
    )

    assert result[["ci_lower", "ci_upper"]].notna().all().all()
    assert (result["ci_lower"] <= result["ci_upper"]).all()


def test_abc_uses_wider_simultaneous_ci_than_ab():
    metrics = _metrics_config().loc[["average"]]
    ab_result = run_ab_test(
        _experiment_data(("A", "B")), metrics, "2026-01-01", "2026-01-07"
    ).iloc[0]
    abc_result = run_abc_test(
        _experiment_data(("A", "B", "C")), metrics, "2026-01-01", "2026-01-07"
    )
    abc_ab = abc_result.loc[abc_result["pair"] == "A_vs_B"].iloc[0]

    ab_width = ab_result["ci_upper"] - ab_result["ci_lower"]
    abc_width = abc_ab["ci_upper"] - abc_ab["ci_lower"]

    assert abc_width > ab_width


def test_abc_returns_simultaneous_ci_for_every_metric_type():
    result = run_abc_test(
        _experiment_data(("A", "B", "C")),
        _metrics_config(),
        "2026-01-01",
        "2026-01-07",
        ci_bootstrap_resamples=500,
        ci_random_state=1,
    )

    assert result[["ci_lower", "ci_upper"]].notna().all().all()
    assert (result["ci_lower"] <= result["ci_upper"]).all()


def test_ab_excel_places_ci_columns_after_power_now(tmp_path: Path):
    result = run_ab_test(
        _experiment_data(("A", "B")),
        _metrics_config().loc[["average"]],
        "2026-01-01",
        "2026-01-07",
    )
    output_path = tmp_path / "ab.xlsx"

    export_ab_results_to_excel(result, output_path)

    worksheet = load_workbook(output_path)["ab_metrics"]
    headers = [cell.value for cell in worksheet[3]]
    power_index = headers.index("power_now")
    assert headers[power_index + 1 : power_index + 3] == ["ci_lower", "ci_upper"]
    assert worksheet.cell(4, power_index + 2).number_format == "#,##0.0000"


def test_ab_keeps_full_precision_p_value_for_statistical_decision():
    control = pd.DataFrame({"value": [1.0, 2.0, 5.0, 8.0]})
    experiment = pd.DataFrame({"value": [3.0, 6.0, 7.0, 12.0, 15.0]})

    actual = calculate_p_value(control, experiment, "average")
    expected = CompareMeans(
        DescrStatsW(control["value"]), DescrStatsW(experiment["value"])
    ).ttest_ind(alternative="two-sided", usevar="unequal")[1]

    assert actual == expected


def test_ratio_point_estimate_uses_the_same_paired_rows_as_its_ci():
    data = pd.DataFrame(
        {
            "user_id": range(6),
            "exp_group": ["A", "A", "A", "B", "B", "B"],
            "num": [2.0, np.nan, 9.0, 4.0, 10.0, np.nan],
            "den": [4.0, 10.0, np.nan, 8.0, 20.0, 40.0],
        }
    )

    control, experiment = _prep_metric_pair_frames(data, metric_type="ratio")
    base, exp, delta, _ = calculate_base_exp_values(control, experiment, "ratio")

    assert base == 0.5
    assert exp == 0.5
    assert delta == 0.0


def test_share_ignores_rows_with_missing_outcome():
    data = pd.DataFrame(
        {
            "user_id": range(4),
            "exp_group": ["A", "A", "B", "B"],
            "success": [1.0, np.nan, 1.0, 0.0],
            "eligible": [1.0, 1.0, 1.0, 1.0],
        }
    )
    metrics = pd.DataFrame(
        [["Share", "share", "success", "eligible", "positive"]],
        index=["share"],
        columns=["desc", "type", "num", "den", "direction"],
    )

    result = run_ab_test(data, metrics, "2026-01-01", "2026-01-07").iloc[0]

    assert result["value_base"] == 1.0
    assert result["value_exp"] == 0.5


def test_ab_reports_missing_input_columns():
    data = pd.DataFrame({"user_id": [1, 2], "exp_group": ["A", "B"]})
    metrics = pd.DataFrame(
        [["Shows", "additive", "shows", "shows", "positive"]],
        index=["shows"],
        columns=["desc", "type", "num", "den", "direction"],
    )

    with pytest.raises(ValueError, match="missing required columns"):
        run_ab_test(data, metrics, "2026-01-01", "2026-01-07")
